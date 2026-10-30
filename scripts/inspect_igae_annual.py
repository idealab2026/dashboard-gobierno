#!/usr/bin/env python3
"""Read official IGAE annual S.13 XLSX strictly as evidence; no Sankey data rewrites.

No inference from Excel cell position may go to the chart until formulas,
headers, vintages and all official economic / COFOG hierarchies are mapped.
"""
import datetime as dt
import hashlib
import io
import json
import urllib.request
from pathlib import Path
from openpyxl import load_workbook

URL="https://www.igae.pap.hacienda.gob.es/sitios/igae/es-ES/Contabilidad/ContabilidadNacional/Publicaciones/Documents/AAPP_A/A_AAPP.xlsx"
OUT=Path("app/data/igae-annual-status.json")
report={"checked_at_utc":dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "source_url":URL,"institution":"IGAE / Ministerio de Hacienda",
        "dataset":"Operaciones no financieras, sector S.13, serie anual",
        "status":"unavailable","published_into_sankey":False,
        "warning":"Sin desglose COFOG de la misma vintage, no reconstruir un Sankey 2025"}
try:
    req=urllib.request.Request(URL,headers={"User-Agent":"IdeaLabOfficialFinanceMonitor/1.0"})
    with urllib.request.urlopen(req,timeout=30) as r:
        content=r.read(7_000_001)
        if r.status!=200:raise ValueError("HTTP non-200")
    if len(content)>7_000_000:raise ValueError("Workbook over 7 MB limit")
    if not content.startswith(b"PK"):raise ValueError("Official response is not an XLSX archive")
    workbook=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    sheets=[]
    for sh in workbook.worksheets[:8]:
        samples=[];last_year=[]
        for row_num,row in enumerate(sh.iter_rows(max_row=min(sh.max_row or 1200,1500),max_col=min(sh.max_column or 15,15),values_only=True),start=1):
            nonempty=[{"col":col+1,"v":str(v).strip()[:92]} for col,v in enumerate(row) if v is not None]
            if row_num<=10 and nonempty:samples.append({"row":row_num,"cells":nonempty[:8]})
            if any(str(v).strip() in ("2025","2024") for v in row[:7] if v is not None) and len(last_year)<9:
                last_year.append({"row":row_num,"cells":nonempty[:11]})
        sheets.append({"name":sh.title,"row_count":sh.max_row,"column_count":sh.max_column,
                       "first_rows":samples,"year_rows":last_year})
    workbook.close()
    report.update(status="downloaded_for_mapping",sha256=hashlib.sha256(content).hexdigest(),
                  bytes=len(content),sheets=sheets)
    print("IGAE_ANNUAL_XLSX sheets",[(s["name"],s["row_count"],s["column_count"]) for s in sheets])
    for sheet in sheets[:2]:
        for row in sheet["first_rows"][:10]:
            print("IGAE_SAMPLE",sheet["name"],row)
        for row in sheet["year_rows"][:6]:
            print("IGAE_YEAR",sheet["name"],row)
except (OSError,ValueError,TimeoutError,ImportError) as error:
    report.update(status="unavailable",error=type(error).__name__+":"+str(error)[:160])
    print("::warning::Unable to inspect official IGAE annual XLSX: "+report["error"])
OUT.parent.mkdir(parents=True,exist_ok=True)
OUT.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print("IGAE_ANNUAL_STATUS",report["status"])
